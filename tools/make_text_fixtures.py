"""Create small, synthetic text documents for local acceptance checks."""
from pathlib import Path


def create_docx(path: Path):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.9)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    for name in ("Title", "Heading 1", "Heading 2"):
        doc.styles[name].font.color.rgb = RGBColor.from_string("172F39")
    doc.add_paragraph("Folio · Documento de prueba", "Title")
    doc.add_paragraph("Ejemplo sintético. No contiene datos reales ni evidencia científica.")
    doc.add_heading("Lectura y estructura", 1)
    paragraph = doc.add_paragraph("La conversión debe conservar ")
    paragraph.add_run("negrita").bold = True
    paragraph.add_run(", ")
    paragraph.add_run("cursiva").italic = True
    paragraph.add_run(" y caracteres como á, ñ, ± y μ. Marcador de control: 9281.")
    doc.add_paragraph("Conservar el orden de los párrafos.", "List Bullet")
    doc.add_paragraph("Mantener los valores de las tablas.", "List Bullet")
    doc.add_paragraph("Seleccionar un documento.", "List Number")
    doc.add_paragraph("Convertirlo y revisar el resultado.", "List Number")
    doc.add_heading("Tabla de ejemplo", 1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Shading Accent 1"
    for cell, value in zip(table.rows[0].cells, ("Grupo", "Muestra", "Cambio")):
        cell.text = value
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(header)
    for row in (("A", "120", "−2,5"), ("B", "98", "+1,2")):
        for cell, value in zip(table.add_row().cells, row):
            cell.text = value
    doc.add_paragraph("Párrafo posterior a la tabla: el orden debe mantenerse.")
    doc.add_heading("Enlace", 2)
    paragraph = doc.add_paragraph()
    hyperlink = OxmlElement("w:hyperlink")
    rid = paragraph.part.relate_to("https://python-docx.readthedocs.io/", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink.set(qn("r:id"), rid)
    run, text = OxmlElement("w:r"), OxmlElement("w:t")
    text.text = "Documentación de python-docx"
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    doc.save(path)


def main():
    root = Path(__file__).resolve().parents[1] / "examples"
    root.mkdir(exist_ok=True)
    create_docx(root / "sample_document.docx")
    (root / "sample_notes.txt").write_text("Notas de prueba\n\nMarcador: 9281.\nEspañol: á, é, í, ó, ú, ñ.\nValores: −2,5; 120; 98.\n\n# Esta línea es texto literal.\n", encoding="utf-8")
    (root / "sample_markdown.md").write_text("# Ejemplo Markdown\n\nTexto con **negrita** y una lista:\n\n- Primer elemento\n- Segundo elemento\n", encoding="utf-8")
    (root / "sample_data.csv").write_text('Grupo;Muestra;Cambio\nA;120;"−2,5"\nB;98;"+1,2"\n', encoding="utf-8")
    (root / "sample_data.tsv").write_text("Grupo\tMuestra\tCambio\nA\t120\t−2,5\nB\t98\t+1,2\n", encoding="utf-8")
    print(root)


if __name__ == "__main__":
    main()
