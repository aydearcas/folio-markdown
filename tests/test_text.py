"""Actual editable-document conversion, routing, fidelity and failure checks."""
import codecs
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from folio.core import ConversionError, Options, convert_file
from folio.formats import detect_format
from tools.make_text_fixtures import create_docx


class TextTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def convert(self, path, **kwargs):
        result = convert_file(path, self.root / "out", Options(**kwargs))
        md = Path(result["markdown"]).read_text(encoding="utf-8")
        report = json.loads((Path(result["output"]) / "conversion.json").read_text(encoding="utf-8"))
        return result, md, report

    def test_txt_without_pdf_engines_and_with_stale_model_path(self):
        path = self.root / "notes.txt"
        path.write_text("Español ñ; μ ± 9281\n# Literal", encoding="utf-8")
        with patch("folio.core.select_engine", side_effect=AssertionError("PDF engine called")):
            result, md, report = self.convert(path, engine="docling", artifacts_path="/nonexistent/models")
        self.assertEqual(result["engine"], "text")
        self.assertIn("ñ; μ ± 9281", md)
        self.assertIn(r"\# Literal", md)
        self.assertIsNone(report["pages"])
        self.assertEqual(report["warnings"], [])

    def test_encodings_bom_and_windows(self):
        for encoding in ("utf-8-sig", "utf-16", "utf-32", "cp1252"):
            with self.subTest(encoding=encoding):
                path = self.root / (encoding + ".txt")
                path.write_bytes("niño café 9281 €".encode(encoding))
                _, md, report = self.convert(path)
                self.assertIn("niño café 9281 €", md)
                self.assertEqual(bool(report["warnings"]), encoding == "cp1252")

    def test_binary_txt_is_rejected(self):
        path = self.root / "renamed.txt"
        path.write_bytes(b"abc\x00\x01")
        with self.assertRaises(ConversionError) as exc:
            self.convert(path)
        self.assertEqual(exc.exception.code, "binary_text")

    def test_empty_text_is_not_published(self):
        path = self.root / "empty.txt"
        path.write_text(" \n")
        with self.assertRaises(ConversionError):
            self.convert(path)
        self.assertFalse(list((self.root / "out").glob("empty__*")))

    def test_markdown_syntax_preserved(self):
        path = self.root / "existing.md"
        content = "# Title\n\n**bold**\n\n![](image.png)\n"
        path.write_text(content, encoding="utf-8")
        _, md, report = self.convert(path)
        self.assertEqual(content, md)
        self.assertIn({"code": "markdown_assets"}, report["warnings"])

    def test_csv_and_tsv_preserve_quoted_newlines_and_columns(self):
        for ext, delimiter in (("csv", ";"), ("tsv", "\t")):
            path = self.root / ("data." + ext)
            path.write_text(f'Group{delimiter}Value\n"A|B"{delimiter}"-2,5\n9281"\n', encoding="utf-8")
            result, md, _ = self.convert(path)
            self.assertIn(r"A\|B", md)
            self.assertIn("-2,5<br>9281", md)
            data = json.loads((Path(result["output"]) / "structure.json").read_text())
            self.assertEqual(data["rows"][1], ["A|B", "-2,5\n9281"])

    def test_forced_types_reject_mismatch(self):
        for filename, content, document_type in (("a.txt", b"hello", "pdf"), ("a.pdf", b"%PDF-1.7", "text")):
            path = self.root / filename
            path.write_bytes(content)
            with self.assertRaises(ConversionError) as exc:
                self.convert(path, document_type=document_type)
            self.assertEqual(exc.exception.code, "type_mismatch")

    def test_signature_before_extension(self):
        path = self.root / "renamed.txt"
        path.write_bytes(b"%PDF-1.7\n")
        self.assertEqual(detect_format(path), "pdf")
        create_docx(path)
        self.assertEqual(detect_format(path), "docx")
        path.write_text("PDF files begin with %PDF-1.7", encoding="utf-8")
        self.assertEqual(detect_format(path), "txt")

    def test_plain_numbered_line_is_literal(self):
        path = self.root / "numbers.txt"
        path.write_text("7. A literal numbered line", encoding="utf-8")
        _, md, _ = self.convert(path)
        self.assertIn(r"7\. A literal", md)

    def test_footnote_and_headers_preserve_content(self):
        from docx.opc.part import Part
        from docx.opc.packuri import PackURI
        doc = Document()
        doc.sections[0].header.paragraphs[0].text = "Running header"
        p = doc.add_paragraph("Text with note")
        ref = OxmlElement("w:footnoteReference")
        ref.set(qn("w:id"), "1")
        p.add_run()._r.append(ref)
        data = b'<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:footnote w:id="1"><w:p><w:r><w:t>Footnote value 42</w:t></w:r></w:p></w:footnote></w:footnotes>'
        part = Part(PackURI("/word/footnotes.xml"), "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml", data, doc.part.package)
        doc.part.relate_to(part, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes")
        path = self.root / "notes.docx"
        doc.save(path)
        _, md, report = self.convert(path, keep_headers=True)
        self.assertIn("Text with note[^fn1]", md)
        self.assertIn("[^fn1]: Footnote value 42", md)
        self.assertIn("Running header", md)
        self.assertIn({"code": "docx_headers_appended"}, report["warnings"])

    def test_docx_structure_and_numbers_in_original_order(self):
        path = self.root / "sample.docx"
        create_docx(path)
        result, md, report = self.convert(path)
        self.assertEqual(result["engine"], "python-docx")
        self.assertIn("# Folio", md)
        self.assertIn("**negrita**", md)
        self.assertIn("*cursiva*", md)
        self.assertIn("- Conservar", md)
        self.assertIn("1. Seleccionar", md)
        self.assertIn("2. Convertirlo", md)
        self.assertIn("| A | 120 | −2,5 |", md)
        self.assertLess(md.index("Tabla de ejemplo"), md.index("| A |"))
        self.assertLess(md.index("| A |"), md.index("Párrafo posterior"))
        self.assertIn("https://python-docx.readthedocs.io/", md)
        self.assertIsNone(report["pages"])
        self.assertEqual(report["warnings"], [])

    def test_docx_merged_cells_have_notice(self):
        path = self.root / "merged.docx"
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).merge(table.cell(0, 1)).text = "Both"
        table.cell(1, 0).text, table.cell(1, 1).text = "12", "34"
        doc.save(path)
        _, md, report = self.convert(path)
        self.assertIn("12 | 34", md)
        self.assertIn({"code": "docx_merged_cells"}, report["warnings"])

    def test_docx_unsupported_content_has_notice(self):
        path = self.root / "tracked.docx"
        doc = Document()
        doc.add_paragraph("Visible paragraph")
        insertion = OxmlElement("w:ins")
        insertion.set(qn("w:id"), "1")
        doc.element.body.append(insertion)
        doc.save(path)
        _, _, report = self.convert(path)
        self.assertTrue(any(w["code"] == "docx_complex" for w in report["warnings"]))

    def test_inline_image_export_and_disabled_notice(self):
        # Valid 1x1 PNG, no imaging dependency required for this fixture.
        import base64
        raw = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII=")
        doc = Document()
        doc.add_paragraph("Inline image test")
        doc.add_picture(io.BytesIO(raw))
        path = self.root / "image.docx"
        doc.save(path)
        result, md, _ = self.convert(path, images=True)
        self.assertIn("images/image_001.png", md)
        self.assertEqual((Path(result["output"]) / "images/image_001.png").read_bytes(), raw)
        _, _, report = self.convert(path, images=False)
        self.assertIn({"code": "docx_images_omitted"}, report["warnings"])

    def test_text_page_copy_is_not_fabricated(self):
        path = self.root / "notes.txt"
        path.write_text("hello")
        result, _, report = self.convert(path, page_copy=True)
        self.assertIn({"code": "text_no_pages"}, report["warnings"])
        self.assertFalse((Path(result["output"]) / "document.pages.md").exists())

    def test_invalid_and_unsupported_files(self):
        for ext, code in (("docx", "invalid_docx"), ("doc", "unsupported_format"), ("rtf", "unsupported_format")):
            path = self.root / ("bad." + ext)
            path.write_bytes(b"not a document")
            with self.assertRaises(ConversionError) as exc:
                self.convert(path)
            self.assertEqual(exc.exception.code, code)


if __name__ == "__main__":
    unittest.main()
